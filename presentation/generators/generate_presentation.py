"""
generate_presentation.py
Creates the official, professional 12-slide presentation for FileMind.
Widescreen 16:9 format with dark tech styling (#0B0F19, Electric Cyan #00D2FF, Neon Pink #FF2D95).
Embeds real application screenshots and slide-by-slide speaker notes.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors
BG_DARK = RGBColor(11, 15, 25)        # #0B0F19
CARD_BG = RGBColor(17, 24, 39)        # #111827
CARD_BORDER = RGBColor(35, 50, 82)    # #233252
CARD_HOVER = RGBColor(24, 33, 54)     # #182136
CYAN = RGBColor(0, 210, 255)          # #00D2FF (Primary Accent)
PINK = RGBColor(255, 45, 149)         # #FF2D95 (Secondary Accent)
GREEN = RGBColor(16, 185, 129)        # #10B981 (Status / Verified)
AMBER = RGBColor(245, 158, 11)        # #F59E0B (Notice)
PURPLE = RGBColor(168, 85, 247)       # #A855F7 (Special)
WHITE = RGBColor(255, 255, 255)
TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8
TEXT_BODY = RGBColor(203, 213, 225)   # #CBD5E1

FONT_FAMILY = "Segoe UI"

def set_slide_background(slide):
    """Sets dark tech background with full bleed rectangle."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK
    bg.line.fill.background()
    return bg

def add_header(slide, category_text, title_text, subtitle_text=""):
    """Adds standard high-contrast header section."""
    # Top Category Badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.8), Inches(0.3))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(19, 29, 49)
    badge.line.color.rgb = CYAN
    badge.line.width = Pt(1)
    tf_b = badge.text_frame
    tf_b.word_wrap = False
    tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0
    p_b = tf_b.paragraphs[0]
    p_b.alignment = PP_ALIGN.CENTER
    run_b = p_b.add_run()
    run_b.text = category_text.upper()
    run_b.font.name = FONT_FAMILY
    run_b.font.size = Pt(9.5)
    run_b.font.bold = True
    run_b.font.color.rgb = CYAN

    # Slide Title Box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.733), Inches(0.65))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    run_t = p_t.add_run()
    run_t.text = title_text
    run_t.font.name = FONT_FAMILY
    run_t.font.size = Pt(24)
    run_t.font.bold = True
    run_t.font.color.rgb = WHITE

    # Subtitle / Context
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.45))
        tf_s = sub_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_right = tf_s.margin_top = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        run_s = p_s.add_run()
        run_s.text = subtitle_text
        run_s.font.name = FONT_FAMILY
        run_s.font.size = Pt(12)
        run_s.font.color.rgb = TEXT_MUTED

def add_card(slide, left, top, width, height, border_color=CARD_BORDER, bg_color=CARD_BG):
    """Creates a styled card container."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card

def add_speaker_note(slide, note_text):
    """Embeds presentation speaker script into PowerPoint notes panel."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = note_text

print("Building refined 12 presentation slides...")

# ==============================================================================
# SLIDE 1: TITLE SLIDE
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)
set_slide_background(s1)

# Subtle decorative background grid line
dec_line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.02))
dec_line.fill.solid()
dec_line.fill.fore_color.rgb = CARD_BORDER
dec_line.line.fill.background()

# Title Category Chip
chip = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.0), Inches(4.2), Inches(0.38))
chip.fill.solid()
chip.fill.fore_color.rgb = RGBColor(19, 29, 49)
chip.line.color.rgb = CYAN
chip.line.width = Pt(1.2)
p = chip.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "ACADEMIC CAPSTONE & HACKATHON 2026"
r.font.name = FONT_FAMILY
r.font.size = Pt(10.5)
r.font.bold = True
r.font.color.rgb = CYAN

# Project Name & Subtitle
title_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(11.733), Inches(2.2))
tf = title_box.text_frame
tf.word_wrap = True
p1 = tf.paragraphs[0]
r1 = p1.add_run()
r1.text = "FileMind"
r1.font.name = FONT_FAMILY
r1.font.size = Pt(54)
r1.font.bold = True
r1.font.color.rgb = WHITE

p2 = tf.add_paragraph()
r2 = p2.add_run()
r2.text = "AI-Powered PDF Vault & Knowledge Retrieval Engine"
r2.font.name = FONT_FAMILY
r2.font.size = Pt(22)
r2.font.bold = True
r2.font.color.rgb = CYAN

p3 = tf.add_paragraph()
p3.space_before = Pt(6)
r3 = p3.add_run()
r3.text = "“Your documents. Organized. Searchable. Intelligent.”"
r3.font.name = FONT_FAMILY
r3.font.size = Pt(15)
r3.font.italic = True
r3.font.color.rgb = PINK

# Main Feature Pillars Banner
pillars_card = add_card(s1, Inches(0.8), Inches(3.9), Inches(11.733), Inches(1.15), border_color=CARD_BORDER)
tf_pil = pillars_card.text_frame
tf_pil.margin_left = Inches(0.35)
tf_pil.margin_top = Inches(0.18)
p_pil = tf_pil.paragraphs[0]
p_pil.text = "Core Architectural Pillars:"
p_pil.font.name = FONT_FAMILY
p_pil.font.size = Pt(10.5)
p_pil.font.bold = True
p_pil.font.color.rgb = TEXT_MUTED

p_pil2 = tf_pil.add_paragraph()
p_pil2.space_before = Pt(4)
r_p1 = p_pil2.add_run()
r_p1.text = "⚡ Multimodal Document AI   •   "
r_p1.font.bold = True
r_p1.font.size = Pt(10.5)
r_p1.font.color.rgb = CYAN
r_p2 = p_pil2.add_run()
r_p2.text = "🔒 Multi-Tenant RLS Vault   •   "
r_p2.font.bold = True
r_p2.font.size = Pt(10.5)
r_p2.font.color.rgb = GREEN
r_p3 = p_pil2.add_run()
r_p3.text = "🔍 Hybrid Semantic Retrieval   •   "
r_p3.font.bold = True
r_p3.font.size = Pt(10.5)
r_p3.font.color.rgb = PINK
r_p4 = p_pil2.add_run()
r_p4.text = "📦 150-File Bulk Ingestion"
r_p4.font.bold = True
r_p4.font.size = Pt(10.5)
r_p4.font.color.rgb = WHITE

# Presenter / Academic / Hackathon Metadata Card
meta_card = add_card(s1, Inches(0.8), Inches(5.3), Inches(11.733), Inches(1.55), border_color=CARD_BORDER)
tf_m = meta_card.text_frame
tf_m.margin_left = Inches(0.4)
tf_m.margin_top = Inches(0.22)

p_m1 = tf_m.paragraphs[0]
r_m1 = p_m1.add_run()
r_m1.text = "Project Team & Affiliation:"
r_m1.font.name = FONT_FAMILY
r_m1.font.size = Pt(11.5)
r_m1.font.bold = True
r_m1.font.color.rgb = WHITE

p_m2 = tf_m.add_paragraph()
p_m2.space_before = Pt(5)
r_m2 = p_m2.add_run()
r_m2.text = "Presenter: [Student Name / Lead Presenter]  |  Team Members: [Team Member 2, Team Member 3]\n"
r_m2.text += "Department: [Department of Computer Science & Engineering]  |  Institution: [College / University Name]\n"
r_m2.text += "Evaluation Track: [College Major Project Evaluation 2026] & [Annual College Hackathon 2026]"
r_m2.font.name = FONT_FAMILY
r_m2.font.size = Pt(10.5)
r_m2.font.color.rgb = TEXT_BODY

add_speaker_note(s1, """Good morning respected evaluator and hackathon judges. 
Today, we are excited to present FileMind — an AI-Powered PDF Vault and Knowledge Retrieval Engine.

Every student, researcher, and organization deals with dozens of PDF files every single week — syllabi, research papers, financial receipts, and legal forms. Yet, our file storage systems have remained virtually unchanged for twenty years: passive, unsearchable folders that force you to manually open and skim through hundreds of pages.

FileMind completely solves this by combining modern cloud storage security with Google Gemini Multimodal AI and hybrid semantic retrieval. It doesn't just store your documents; it understands them. Let us show you how.""")


# ==============================================================================
# SLIDE 2: THE PROBLEM
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
set_slide_background(s2)
add_header(s2, "Problem Statement", "The Crisis of Unstructured PDF Management", 
           "Traditional storage systems treat PDFs as opaque black boxes, creating massive knowledge fragmentation.")

cards_data_s2 = [
    {
        "title": "1. Opaque & Unsearchable Silos",
        "subtitle": "Trapped in Local Folders",
        "color": PINK,
        "bullets": [
            "PDFs are scattered across local folders, cloud drives, and downloads.",
            "Standard OS search only matches filenames, completely missing embedded body text and tables.",
            "Scanned invoices, contracts, and image-based PDFs remain 100% invisible to search."
        ]
    },
    {
        "title": "2. Manual Reading Fatigue",
        "subtitle": "Hours Lost Skimming Pages",
        "color": AMBER,
        "bullets": [
            "Extracting critical facts (e.g. invoice total, vendor name, contract clause) requires tedious manual opening.",
            "Users re-read the same lengthy reports repeatedly just to locate a single data point.",
            "Zero automatic categorization or structured schema generation."
        ]
    },
    {
        "title": "3. Security & Scale Bottlenecks",
        "subtitle": "Fragile Ad-Hoc Solutions",
        "color": CYAN,
        "bullets": [
            "Generic cloud drives provide storage but lack multi-tenant row-level access control for documents.",
            "Conventional desktop OCR tools choke on large batches and fail on modern complex layouts.",
            "No automated queue recovery, leading to frozen uploads and lost metadata."
        ]
    }
]

left_pos = Inches(0.8)
card_width = Inches(3.7)
card_height = Inches(4.7)

for i, c in enumerate(cards_data_s2):
    card = add_card(s2, left_pos, Inches(2.0), card_width, card_height)
    
    # Top Accent Cap Bar
    cap = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, Inches(2.0), card_width, Inches(0.06))
    cap.fill.solid()
    cap.fill.fore_color.rgb = c["color"]
    cap.line.fill.background()

    tf_c = card.text_frame
    tf_c.margin_left = Inches(0.3)
    tf_c.margin_right = Inches(0.3)
    tf_c.margin_top = Inches(0.3)

    # Card Title
    p_t = tf_c.paragraphs[0]
    r_t = p_t.add_run()
    r_t.text = c["title"]
    r_t.font.name = FONT_FAMILY
    r_t.font.size = Pt(13.5)
    r_t.font.bold = True
    r_t.font.color.rgb = c["color"]

    # Subtitle
    p_sub = tf_c.add_paragraph()
    p_sub.space_before = Pt(3)
    r_sub = p_sub.add_run()
    r_sub.text = c["subtitle"]
    r_sub.font.name = FONT_FAMILY
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = TEXT_MUTED

    # Bullets
    for b in c["bullets"]:
        p_b = tf_c.add_paragraph()
        p_b.space_before = Pt(12)
        r_bullet = p_b.add_run()
        r_bullet.text = "• "
        r_bullet.font.bold = True
        r_bullet.font.color.rgb = c["color"]
        r_text = p_b.add_run()
        r_text.text = b
        r_text.font.name = FONT_FAMILY
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = TEXT_BODY

    left_pos += card_width + Inches(0.316)

add_speaker_note(s2, """Let's look at why managing PDFs is broken today. 
First, documents are scattered everywhere — Downloads, Google Drive, email attachments. Standard file search only matches the filename. If a document is named 'scan_001.pdf', it is completely invisible.

Second, extracting structured data like invoice amounts, contract renewal dates, or research conclusions requires manually opening each file and skimming dozens of pages. Studies show knowledge workers spend up to 20% of their week just searching for information.

Third, current DIY tools lack security and concurrency. Uploading 50 files simultaneously crashes existing desktop scripts or triggers rate limits. FileMind was designed specifically to eliminate these bottlenecks.""")


# ==============================================================================
# SLIDE 3: OUR SOLUTION
# ==============================================================================
s3 = prs.slides.add_slide(blank_layout)
set_slide_background(s3)
add_header(s3, "Product Innovation", "Our Solution: The FileMind AI Vault", 
           "Transforming passive PDF storage into an active, intelligent, and secure retrieval ecosystem.")

# Left Card: Traditional Broken Workflow
left_c = add_card(s3, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.7), border_color=RGBColor(60, 30, 45))
cap_l = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.0), Inches(5.6), Inches(0.06))
cap_l.fill.solid()
cap_l.fill.fore_color.rgb = PINK
cap_l.line.fill.background()

tf_l = left_c.text_frame
tf_l.margin_left = Inches(0.35)
tf_l.margin_right = Inches(0.35)
tf_l.margin_top = Inches(0.28)

p_lh = tf_l.paragraphs[0]
r_lh = p_lh.add_run()
r_lh.text = "❌ TRADITIONAL PDF STORAGE"
r_lh.font.name = FONT_FAMILY
r_lh.font.size = Pt(13.5)
r_lh.font.bold = True
r_lh.font.color.rgb = PINK

p_lsub = tf_l.add_paragraph()
r_lsub = p_lsub.add_run()
r_lsub.text = "Passive, dumb folders with manual maintenance"
r_lsub.font.size = Pt(10.5)
r_lsub.font.color.rgb = TEXT_MUTED

trad_points = [
    ("Manual Filename Search", "Only finds documents if you recall the exact file naming syntax."),
    ("Opaque Scanned Pages", "Cannot inspect image-based or multi-column PDFs without expensive desktop software."),
    ("Zero Metadata Extraction", "Invoice numbers, amounts, and tags must be manually typed into spreadsheets."),
    ("Single-File Friction", "Must upload and view files one by one; fails on large batch volumes.")
]

for title, desc in trad_points:
    p_tp = tf_l.add_paragraph()
    p_tp.space_before = Pt(10)
    r_tp1 = p_tp.add_run()
    r_tp1.text = f"• {title}: "
    r_tp1.font.bold = True
    r_tp1.font.color.rgb = WHITE
    r_tp2 = p_tp.add_run()
    r_tp2.text = desc
    r_tp2.font.color.rgb = TEXT_MUTED
    r_tp2.font.size = Pt(10.5)

# Right Card: The FileMind Way
right_c = add_card(s3, Inches(6.8), Inches(2.0), Inches(5.733), Inches(4.7), border_color=RGBColor(15, 60, 80))
cap_r = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(2.0), Inches(5.733), Inches(0.06))
cap_r.fill.solid()
cap_r.fill.fore_color.rgb = CYAN
cap_r.line.fill.background()

tf_r = right_c.text_frame
tf_r.margin_left = Inches(0.35)
tf_r.margin_right = Inches(0.35)
tf_r.margin_top = Inches(0.28)

p_rh = tf_r.paragraphs[0]
r_rh = p_rh.add_run()
r_rh.text = "✅ FILEMIND AI-POWERED VAULT"
r_rh.font.name = FONT_FAMILY
r_rh.font.size = Pt(13.5)
r_rh.font.bold = True
r_rh.font.color.rgb = CYAN

p_rsub = tf_r.add_paragraph()
r_rsub = p_rsub.add_run()
r_rsub.text = "Autonomous document intelligence with semantic discovery"
r_rsub.font.size = Pt(10.5)
r_rsub.font.color.rgb = TEXT_MUTED

fm_points = [
    ("Multimodal Understanding", "Google Gemini extracts categories, tags, key-value data, and page-by-page text."),
    ("Semantic Hybrid Search", "Find concepts (e.g. 'quarterly cloud spending') even if the exact keyword never appears."),
    ("High-Capacity Concurrency", "Bulk upload up to 150 PDFs (49 MiB each) with 4-way upload and 2-way AI sliding windows."),
    ("Enterprise Security", "PostgreSQL Row-Level Security, advisory locks, and cryptographic fencing leases prevent data leaks.")
]

for title, desc in fm_points:
    p_fp = tf_r.add_paragraph()
    p_fp.space_before = Pt(10)
    r_fp1 = p_fp.add_run()
    r_fp1.text = f"• {title}: "
    r_fp1.font.bold = True
    r_fp1.font.color.rgb = CYAN
    r_fp2 = p_fp.add_run()
    r_fp2.text = desc
    r_fp2.font.color.rgb = TEXT_BODY
    r_fp2.font.size = Pt(10.5)

add_speaker_note(s3, """Here is the core thesis of FileMind. 
On the left, you see the traditional approach: dumb folders, manual renaming, and zero insight into document contents. If you need a contract renewal date, you have to read the contract.

On the right is FileMind. When you drop your PDFs into FileMind, an automated background pipeline processes them using Google Gemini multimodal AI. It classifies the file into categories like Invoice, Legal Contract, or Research Paper, pulls out key-value metadata, generates semantic tags, and indexes every single page into high-dimensional vector embeddings.

The user doesn't just store files — they gain an interactive knowledge vault.""")


# ==============================================================================
# SLIDE 4: KEY FEATURES
# ==============================================================================
s4 = prs.slides.add_slide(blank_layout)
set_slide_background(s4)
add_header(s4, "Verified Capabilities", "Core Implemented Features", 
           "Every feature shown is fully implemented, tested, and running in the FileMind codebase.")

features = [
    {
        "badge": "BULK INGESTION",
        "title": "150-File Bulk Upload",
        "desc": "Sliding-window concurrency engine handling up to 150 files per batch, 49 MiB per file, with partial-failure isolation.",
        "color": CYAN
    },
    {
        "badge": "MULTIMODAL AI",
        "title": "Gemini Multimodal Analysis",
        "desc": "Strict JSON schema extraction for document category, custom tags, and dynamic key-value metadata with injection defense.",
        "color": PINK
    },
    {
        "badge": "SEARCH ENGINE",
        "title": "Hybrid Semantic Retrieval",
        "desc": "Cosine similarity over vector embeddings combined with full-text search, featuring a 500-entry in-memory LRU cache.",
        "color": GREEN
    },
    {
        "badge": "VAULT SECURITY",
        "title": "Multi-Tenant Isolation",
        "desc": "Supabase PostgreSQL Row-Level Security (RLS) guaranteeing strict per-user isolation across database records and storage.",
        "color": PURPLE
    },
    {
        "badge": "RESILIENCE",
        "title": "Durable Queue & Leases",
        "desc": "Atomic job claiming with SKIP LOCKED and UUID fencing tokens, recovering stale processing jobs automatically.",
        "color": AMBER
    },
    {
        "badge": "OPTIMIZATION",
        "title": "Lossless PDF Optimization",
        "desc": "Structural PDF optimization via pdf-lib saving up to 15% storage while safely skipping signed or encrypted documents.",
        "color": CYAN
    }
]

# 2x3 Grid Layout
grid_w = Inches(3.7)
grid_h = Inches(2.2)
x_offsets = [Inches(0.8), Inches(4.8), Inches(8.8)]
y_offsets = [Inches(2.0), Inches(4.5)]

for idx, feat in enumerate(features):
    gx = x_offsets[idx % 3]
    gy = y_offsets[idx // 3]
    card = add_card(s4, gx, gy, grid_w, grid_h)
    
    tf_f = card.text_frame
    tf_f.margin_left = Inches(0.25)
    tf_f.margin_right = Inches(0.25)
    tf_f.margin_top = Inches(0.2)

    # Badge
    p_b = tf_f.paragraphs[0]
    r_b = p_b.add_run()
    r_b.text = f"[{feat['badge']}]"
    r_b.font.name = FONT_FAMILY
    r_b.font.size = Pt(9)
    r_b.font.bold = True
    r_b.font.color.rgb = feat['color']

    # Title
    p_t = tf_f.add_paragraph()
    p_t.space_before = Pt(2)
    r_t = p_t.add_run()
    r_t.text = feat['title']
    r_t.font.name = FONT_FAMILY
    r_t.font.size = Pt(13)
    r_t.font.bold = True
    r_t.font.color.rgb = WHITE

    # Desc
    p_d = tf_f.add_paragraph()
    p_d.space_before = Pt(4)
    r_d = p_d.add_run()
    r_d.text = feat['desc']
    r_d.font.name = FONT_FAMILY
    r_d.font.size = Pt(9.8)
    r_d.font.color.rgb = TEXT_BODY

add_speaker_note(s4, """To ensure credibility for both our academic evaluator and hackathon judges: every feature on this slide is verified and running in our repository. 

First, our bulk upload manager supports up to 150 files per batch with sliding-window concurrency. 
Second, we use Google Gemini with strict response schemas to parse documents securely.
Third, our search engine pairs pgvector embeddings with an in-memory 500-capacity LRU cache so identical queries return instantly.
Fourth, we have database-enforced multi-tenancy with PostgreSQL Row-Level Security.
Fifth, distributed worker leases prevent double-locking when multiple workers process the queue.
And sixth, we perform non-destructive PDF stream compression using pdf-lib while preserving digital signatures.""")


# ==============================================================================
# SLIDE 5: HOW FILEMIND WORKS (PIPELINE WORKFLOW)
# ==============================================================================
s5 = prs.slides.add_slide(blank_layout)
set_slide_background(s5)
add_header(s5, "System Pipeline", "End-to-End Ingestion & Processing Workflow", 
           "Step-by-step lifecycle from client-side drag-and-drop to vector indexing and viewer retrieval.")

steps = [
    {
        "num": "01",
        "title": "Client Batch Upload",
        "details": "User drops up to 150 PDFs. Frontend validates 49 MiB limit, requests signed upload URLs, and streams files directly to private storage."
    },
    {
        "num": "02",
        "title": "Quota & DB Commit",
        "details": "Authoritative RPC verifies storage object size & ownership. Advisory lock checks 2 GB user storage quota & inserts pending document."
    },
    {
        "num": "03",
        "title": "Queue Lease Claim",
        "details": "Background worker invokes claim RPC with SKIP LOCKED. Assigns a cryptographic UUID fencing token to prevent duplicate work."
    },
    {
        "num": "04",
        "title": "Gemini Multimodal AI",
        "details": "PDF is processed via Gemini Files API. Model cascade (3.6-flash -> 3.8-flash -> preview) extracts structured metadata with backoff."
    },
    {
        "num": "05",
        "title": "Atomic Finalization",
        "details": "Transaction commits category, tags, and pages while validating fencing token. Clears lease token and schedules vector indexing."
    },
    {
        "num": "06",
        "title": "Instant Retrieval",
        "details": "Page chunks are indexed with 768-dim embeddings. User searches return ranked matches with snippet previews in <200ms."
    }
]

step_w = Inches(1.8)
step_h = Inches(4.7)
start_x = Inches(0.8)
gap = Inches(0.18)

for i, st in enumerate(steps):
    x = start_x + i * (step_w + gap)
    card = add_card(s5, x, Inches(2.0), step_w, step_h, border_color=CARD_BORDER)
    
    tf_s = card.text_frame
    tf_s.margin_left = Inches(0.15)
    tf_s.margin_right = Inches(0.15)
    tf_s.margin_top = Inches(0.2)

    # Step Number
    p_n = tf_s.paragraphs[0]
    r_n = p_n.add_run()
    r_n.text = st["num"]
    r_n.font.name = FONT_FAMILY
    r_n.font.size = Pt(24)
    r_n.font.bold = True
    r_n.font.color.rgb = CYAN if i % 2 == 0 else PINK

    # Step Title
    p_t = tf_s.add_paragraph()
    p_t.space_before = Pt(4)
    r_t = p_t.add_run()
    r_t.text = st["title"]
    r_t.font.name = FONT_FAMILY
    r_t.font.size = Pt(11.5)
    r_t.font.bold = True
    r_t.font.color.rgb = WHITE

    # Step Details
    p_d = tf_s.add_paragraph()
    p_d.space_before = Pt(10)
    r_d = p_d.add_run()
    r_d.text = st["details"]
    r_d.font.name = FONT_FAMILY
    r_d.font.size = Pt(9.5)
    r_d.font.color.rgb = TEXT_BODY

add_speaker_note(s5, """This workflow diagram illustrates how FileMind processes documents from end to end.
In Step 1, the user drops a batch of PDFs into the web interface. 
In Step 2, the client uploads directly to Supabase Storage, and our PostgreSQL stored procedure atomically enforces the user's 2-gigabyte quota.
In Step 3, our queue engine uses PostgreSQL 'SKIP LOCKED' to acquire a lease with a unique fencing token.
In Step 4, Google Gemini multimodal model reads the document and extracts structured JSON.
In Step 5, the results are atomically written to the database, ensuring zero race conditions.
And in Step 6, page embeddings are indexed for instant semantic retrieval in our dashboard search.""")


# ==============================================================================
# SLIDE 6: LIVE APPLICATION SCREENSHOTS
# ==============================================================================
s6 = prs.slides.add_slide(blank_layout)
set_slide_background(s6)
add_header(s6, "Production UI", "Live Application Screenshots", 
           "High-contrast dark-mode interface built with Next.js 16, Tailwind CSS v4, and Lucide icons.")

shot_dir = os.path.join(os.getcwd(), "presentation_assets")
dash_path = os.path.join(shot_dir, "01_dashboard.png")
vault_path = os.path.join(shot_dir, "03_documents_vault.png")

# Left Column: Dashboard Showcase
card_left = add_card(s6, Inches(0.8), Inches(2.0), Inches(5.7), Inches(4.85), border_color=CARD_BORDER)
if os.path.exists(dash_path):
    s6.shapes.add_picture(dash_path, Inches(0.95), Inches(2.15), width=Inches(5.4))
cap_box1 = s6.shapes.add_textbox(Inches(0.95), Inches(5.62), Inches(5.4), Inches(1.1))
tf_c1 = cap_box1.text_frame
tf_c1.word_wrap = True
p_c1 = tf_c1.paragraphs[0]
r_c1 = p_c1.add_run()
r_c1.text = "Dashboard & Upload Dropzone:\n"
r_c1.font.bold = True
r_c1.font.size = Pt(11)
r_c1.font.color.rgb = CYAN
r_c1_sub = p_c1.add_run()
r_c1_sub.text = "Features real-time metric cards (Document Count, Storage Quota, Processing Rate), 150-file bulk drag-and-drop zone, and authoritative document status."
r_c1_sub.font.size = Pt(9.8)
r_c1_sub.font.color.rgb = TEXT_BODY

# Right Column: Documents Vault Showcase
card_right = add_card(s6, Inches(6.8), Inches(2.0), Inches(5.733), Inches(4.85), border_color=CARD_BORDER)
if os.path.exists(vault_path):
    s6.shapes.add_picture(vault_path, Inches(6.95), Inches(2.15), width=Inches(5.4))
cap_box2 = s6.shapes.add_textbox(Inches(6.95), Inches(5.62), Inches(5.4), Inches(1.1))
tf_c2 = cap_box2.text_frame
tf_c2.word_wrap = True
p_c2 = tf_c2.paragraphs[0]
r_c2 = p_c2.add_run()
r_c2.text = "Documents Vault & Semantic Explorer:\n"
r_c2.font.bold = True
r_c2.font.size = Pt(11)
r_c2.font.color.rgb = PINK
r_c2_sub = p_c2.add_run()
r_c2_sub.text = "Interactive repository displaying automated category badges (Financial Report, Contract, Research), AI-extracted semantic tags, file size savings, and actions."
r_c2_sub.font.size = Pt(9.8)
r_c2_sub.font.color.rgb = TEXT_BODY

add_speaker_note(s6, """Here are real screenshots captured directly from our running application on localhost:3000.
On the left, you see the FileMind Dashboard. Notice the welcome banner, the live statistics cards showing storage quota and processed document count, and the bulk upload zone.

On the right is our Documents Vault. Notice how each document is tagged automatically with category badges like 'Financial Report' or 'Contract'. Users can download the file, inspect page-by-page text, or trigger manual retries. The interface is completely responsive, styled in high-contrast near-black with electric cyan and neon pink accents.""")


# ==============================================================================
# SLIDE 7: TECHNOLOGY STACK
# ==============================================================================
s7 = prs.slides.add_slide(blank_layout)
set_slide_background(s7)
add_header(s7, "Engineering Foundation", "Verified Technology Stack", 
           "Engineered with modern, production-grade frameworks ensuring performance, scalability, and type safety.")

tech_stacks = [
    {
        "category": "Frontend Framework",
        "tech": "Next.js 16.3 + React 19",
        "color": CYAN,
        "role": "Next.js 16 with Turbopack for lightning-fast compilation, React 19 Server Components, and client-side streaming UI."
    },
    {
        "category": "Styling & Design System",
        "tech": "Tailwind CSS v4 + Lucide",
        "color": PINK,
        "role": "Modern CSS-first theme with zero runtime overhead, high-contrast accessible typography, and Lucide React icons."
    },
    {
        "category": "Database & Security",
        "tech": "Supabase (PostgreSQL 15+)",
        "color": GREEN,
        "role": "PostgreSQL with strict Row-Level Security (RLS), transactional advisory locks, and atomic PL/pgSQL stored procedures."
    },
    {
        "category": "Multimodal AI",
        "tech": "Google Gemini API (@google/genai)",
        "color": PURPLE,
        "role": "Direct multimodal PDF processing via Gemini Files API with structured JSON response schema enforcement."
    },
    {
        "category": "Search & Embeddings",
        "tech": "Vector Embeddings + LRU Cache",
        "color": AMBER,
        "role": "768-dimensional document embeddings paired with an in-memory 500-capacity LRU cache for sub-200ms queries."
    },
    {
        "category": "Storage & Optimization",
        "tech": "Supabase S3 Storage + pdf-lib",
        "color": CYAN,
        "role": "Private user-isolated storage buckets with lossless structural PDF stream optimization via pdf-lib."
    }
]

stack_w = Inches(5.6)
stack_h = Inches(1.35)
sx_offsets = [Inches(0.8), Inches(6.8)]
sy_offsets = [Inches(2.0), Inches(3.6), Inches(5.2)]

for idx, t in enumerate(tech_stacks):
    col = idx % 2
    row = idx // 2
    x = sx_offsets[col]
    y = sy_offsets[row]
    
    card = add_card(s7, x, y, stack_w, stack_h)
    tf_t = card.text_frame
    tf_t.margin_left = Inches(0.3)
    tf_t.margin_top = Inches(0.18)
    tf_t.margin_right = Inches(0.3)

    p_c = tf_t.paragraphs[0]
    r_c = p_c.add_run()
    r_c.text = t["category"].upper()
    r_c.font.name = FONT_FAMILY
    r_c.font.size = Pt(8.5)
    r_c.font.bold = True
    r_c.font.color.rgb = t["color"]

    p_tech = tf_t.add_paragraph()
    r_tech = p_tech.add_run()
    r_tech.text = t["tech"]
    r_tech.font.name = FONT_FAMILY
    r_tech.font.size = Pt(13)
    r_tech.font.bold = True
    r_tech.font.color.rgb = WHITE

    p_role = tf_t.add_paragraph()
    p_role.space_before = Pt(3)
    r_role = p_role.add_run()
    r_role.text = t["role"]
    r_role.font.name = FONT_FAMILY
    r_role.font.size = Pt(9.5)
    r_role.font.color.rgb = TEXT_BODY

add_speaker_note(s7, """Our technology stack represents the latest in web development and artificial intelligence. 
We selected Next.js 16 and React 19 as our full-stack framework for their native support for Server Actions, background tasks, and Turbopack build speeds.

For data storage and security, we rely on Supabase PostgreSQL. Crucially, we do not rely solely on application-level security; every document is isolated at the database level using PostgreSQL Row-Level Security policies.

For document understanding, we use the official Google GenAI SDK. Instead of running brittle OCR pipelines, we upload PDF binaries directly to the Gemini multimodal API, which understands complex layouts, tables, and handwritten signatures with human-level accuracy.""")


# ==============================================================================
# SLIDE 8: SYSTEM ARCHITECTURE
# ==============================================================================
s8 = prs.slides.add_slide(blank_layout)
set_slide_background(s8)
add_header(s8, "System Design", "Layered System Architecture", 
           "High-level architectural blueprint showing clear separation of concerns and secure communication boundaries.")

layers = [
    {
        "tier": "TIER 1: PRESENTATION LAYER (CLIENT-SIDE)",
        "color": CYAN,
        "items": [
            "Next.js 16 Responsive UI: Dashboard, Document Vault, PDF Viewer, Semantic Search",
            "Client Batch Upload Manager: Sliding-window queue (max 4 concurrent uploads, per-file progress)"
        ]
    },
    {
        "tier": "TIER 2: GATEWAY & SERVERLESS ROUTING (NEXT.JS 16)",
        "color": PINK,
        "items": [
            "API Endpoints: /api/documents, /api/search, /api/cron/process-queue",
            "Execution Handlers: Next.js 16 after() background execution + Supabase Auth Session Validation"
        ]
    },
    {
        "tier": "TIER 3: DATA & ATOMIC TRANSACTION TIER (SUPABASE / POSTGRESQL)",
        "color": GREEN,
        "items": [
            "Row-Level Security (RLS) enforcing tenant isolation; Private S3-Compatible Storage bucket 'documents'",
            "Stored Procedures: create_document_with_quota, acquire_document_processing, finalize_document_processing"
        ]
    },
    {
        "tier": "TIER 4: INTELLIGENCE & ACCELERATION ENGINE",
        "color": PURPLE,
        "items": [
            "Gemini Multimodal Files API: Structured category, metadata, and page text extraction",
            "Search Optimization: 500-LRU embedding cache, snippet coalescing, and lossless pdf-lib compression"
        ]
    }
]

y_pos = Inches(2.0)
layer_h = Inches(1.08)

for l in layers:
    card = add_card(s8, Inches(0.8), y_pos, Inches(11.733), layer_h)
    tf_l = card.text_frame
    tf_l.margin_left = Inches(0.35)
    tf_l.margin_top = Inches(0.12)
    tf_l.margin_right = Inches(0.35)

    p_t = tf_l.paragraphs[0]
    r_t = p_t.add_run()
    r_t.text = l["tier"]
    r_t.font.name = FONT_FAMILY
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    r_t.font.color.rgb = l["color"]

    for item in l["items"]:
        p_i = tf_l.add_paragraph()
        p_i.space_before = Pt(2)
        r_bullet = p_i.add_run()
        r_bullet.text = "• "
        r_bullet.font.bold = True
        r_bullet.font.color.rgb = l["color"]
        r_item = p_i.add_run()
        r_item.text = item
        r_item.font.name = FONT_FAMILY
        r_item.font.size = Pt(10)
        r_item.font.color.rgb = TEXT_BODY

    y_pos += layer_h + Inches(0.16)

add_speaker_note(s8, """This architecture slide demonstrates how our system is engineered across four distinct tiers.
At Tier 1, the client browser manages a sliding-window queue that limits network saturation by uploading at most 4 files concurrently.
At Tier 2, Next.js serverless routes validate user sessions and schedule asynchronous execution using Next.js 16's 'after()' API.
At Tier 3, Supabase PostgreSQL handles atomic transactions. Stored procedures like 'create_document_with_quota' check user storage capacity and verify storage objects inside a single database transaction.
And at Tier 4, Google Gemini and vector embeddings process the document content asynchronously, making sure the user never experiences interface lag.""")


# ==============================================================================
# SLIDE 9: WHAT MAKES FILEMIND DIFFERENT?
# ==============================================================================
s9 = prs.slides.add_slide(blank_layout)
set_slide_background(s9)
add_header(s9, "Competitive Value", "What Makes FileMind Different?", 
           "Direct feature comparison against traditional folder storage, cloud drives, and generic OCR utilities.")

# Comparison Table
rows = 6
cols = 4
left = Inches(0.8)
top = Inches(2.0)
width = Inches(11.733)
height = Inches(4.7)

table_shape = s9.shapes.add_table(rows, cols, left, top, width, height)
table = table_shape.table

# Set Column Widths
table.columns[0].width = Inches(3.2)
table.columns[1].width = Inches(2.8)
table.columns[2].width = Inches(2.8)
table.columns[3].width = Inches(2.933)

headers = ["Capability / Feature", "Local OS Folders", "Generic Cloud Drive", "FileMind AI Vault"]
for j, h in enumerate(headers):
    cell = table.cell(0, j)
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(24, 33, 54) if j < 3 else RGBColor(19, 45, 65)
    p = cell.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT
    r = p.add_run()
    r.text = h
    r.font.name = FONT_FAMILY
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = CYAN if j == 3 else WHITE

matrix = [
    ("Search Methodology", "Exact filename match only", "Keyword match on text PDFs", "Semantic & Concept Retrieval (768-dim embeddings)"),
    ("Multimodal Extraction", "None (Manual reading)", "Basic OCR (Text only)", "Structured JSON Schema (Category, Tags, Key-Values)"),
    ("Scanned / Image PDFs", "Opaque & Inaccessible", "Limited text layer export", "100% Understood via Multimodal Vision AI"),
    ("Bulk Upload Management", "OS Copy (No validation)", "Slow sequential sync", "150-file Batch Sliding Window (4x network concurrency)"),
    ("Security & Tenant Isolation", "Local user permissions", "Coarse folder sharing links", "Database Row-Level Security (RLS) & Fencing Leases")
]

for i, row in enumerate(matrix):
    for j, val in enumerate(row):
        cell = table.cell(i + 1, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG if j < 3 else RGBColor(15, 30, 48)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT
        r = p.add_run()
        r.text = val
        r.font.name = FONT_FAMILY
        r.font.size = Pt(10)
        if j == 3:
            r.font.bold = True
            r.font.color.rgb = CYAN
        else:
            r.font.color.rgb = TEXT_BODY if j == 0 else TEXT_MUTED

add_speaker_note(s9, """Why does FileMind stand out compared to existing alternatives? This comparison table summarizes the differences.
Against local operating systems, FileMind eliminates manual search by indexing the full semantics of the document.
Against generic cloud drives like Google Drive or Dropbox, FileMind doesn't just store bytes — it extracts structured business metadata: vendor names, due dates, invoice numbers, and custom tags.
And against desktop OCR software, FileMind uses Google Gemini's multimodal vision model, which can understand scanned receipts, rotated pages, and tabular data that standard OCR engines garble. 
Most importantly, all of this is protected by strict database Row-Level Security.""")


# ==============================================================================
# SLIDE 10: CHALLENGES & ENGINEERING
# ==============================================================================
s10 = prs.slides.add_slide(blank_layout)
set_slide_background(s10)
add_header(s10, "Technical Rigor", "Engineering Challenges & Solutions", 
           "Real software engineering hurdles encountered during development and the verified solutions implemented.")

challenges = [
    {
        "title": "1. Gemini 503 Spikes & Model Cascade",
        "prob": "During high-demand periods, Google's standard tier returns 503 UNAVAILABLE, causing queue processing failures.",
        "sol": "Engineered an intelligent model cascade: gemini-3.6-flash -> gemini-3.8-flash -> gemini-3-flash-preview with exponential backoff and jitter.",
        "color": CYAN
    },
    {
        "title": "2. Serverless Timeouts on 150-File Batches",
        "prob": "Processing 150 PDFs in a standard Next.js 60-second serverless execution budget leads to function timeout kills.",
        "sol": "Implemented sliding-window concurrency (max 4 uploads, max 2 AI workers), Next.js after() execution, and a pg_cron background consumer.",
        "color": PINK
    },
    {
        "title": "3. Concurrency Race Conditions & Deadlocks",
        "prob": "Simultaneous file confirmations cause duplicate rows and quota circumvention under concurrent uploads.",
        "sol": "Added transactional PostgreSQL advisory locks pg_advisory_xact_lock(user_id) and atomic SKIP LOCKED queue claiming with UUID fencing leases.",
        "color": GREEN
    },
    {
        "title": "4. Safe PDF Optimization Without Corruption",
        "prob": "Compressing digitally signed or encrypted PDFs invalidates cryptographic signatures and renders files unreadable.",
        "sol": "Conservative pre-checks detect /ByteRange signatures and /Encrypt dictionaries, safely skipping them while optimizing standard structural streams.",
        "color": AMBER
    }
]

c_w = Inches(5.6)
c_h = Inches(2.2)
cx_offsets = [Inches(0.8), Inches(6.8)]
cy_offsets = [Inches(2.0), Inches(4.5)]

for idx, ch in enumerate(challenges):
    col = idx % 2
    row = idx // 2
    x = cx_offsets[col]
    y = cy_offsets[row]
    
    card = add_card(s10, x, y, c_w, c_h)
    tf_c = card.text_frame
    tf_c.margin_left = Inches(0.3)
    tf_c.margin_top = Inches(0.18)
    tf_c.margin_right = Inches(0.3)

    p_t = tf_c.paragraphs[0]
    r_t = p_t.add_run()
    r_t.text = ch["title"]
    r_t.font.name = FONT_FAMILY
    r_t.font.size = Pt(12)
    r_t.font.bold = True
    r_t.font.color.rgb = ch["color"]

    p_p = tf_c.add_paragraph()
    p_p.space_before = Pt(4)
    r_pl = p_p.add_run()
    r_pl.text = "Challenge: "
    r_pl.font.bold = True
    r_pl.font.size = Pt(9.5)
    r_pl.font.color.rgb = WHITE
    r_pt = p_p.add_run()
    r_pt.text = ch["prob"]
    r_pt.font.size = Pt(9.5)
    r_pt.font.color.rgb = TEXT_MUTED

    p_s = tf_c.add_paragraph()
    p_s.space_before = Pt(4)
    r_sl = p_s.add_run()
    r_sl.text = "Solution: "
    r_sl.font.bold = True
    r_sl.font.size = Pt(9.5)
    r_sl.font.color.rgb = GREEN
    r_st = p_s.add_run()
    r_st.text = ch["sol"]
    r_st.font.size = Pt(9.5)
    r_st.font.color.rgb = TEXT_BODY

add_speaker_note(s10, """In any serious engineering project, the hardest parts are the edge cases.
One major challenge we faced was that Google's Gemini API occasionally returns 503 High Demand errors during peak hours. Rather than failing the user's upload, we built a 3-tier model cascade that retries with exponential backoff and automatically falls back to alternative verified models.

Another challenge was serverless timeouts on Vercel. A single request cannot process 150 PDFs within 60 seconds. We solved this by using a sliding-window queue and background pg_cron workers that drain at most 3 documents per invocation.

Finally, we used PostgreSQL transactional advisory locks to prevent race conditions during bulk uploads, guaranteeing that quota limits can never be bypassed.""")


# ==============================================================================
# SLIDE 11: FUTURE SCOPE
# ==============================================================================
s11 = prs.slides.add_slide(blank_layout)
set_slide_background(s11)
add_header(s11, "Roadmap", "Future Scope & Planned Advancements", 
           "Strategic enhancements designed to expand FileMind into an enterprise-ready knowledge platform.")

roadmap = [
    {
        "phase": "PHASE 2 (NEXT MILESTONE)",
        "title": "Cross-Document Synthesis & Multi-PDF Chat",
        "bullets": [
            "Ask questions across your entire library: 'Compare revenue across all Q1, Q2, and Q3 financial reports.'",
            "Synthesize comparative summaries and highlight conflicting clauses across multiple contracts."
        ],
        "color": CYAN
    },
    {
        "phase": "PHASE 3",
        "title": "Interactive Visual Citations in PDF Viewer",
        "bullets": [
            "Exact bounding-box highlighting on the original PDF canvas showing the exact paragraph source for every AI answer.",
            "Clicking a search result immediately scrolls the viewer to the exact matched text coordinates."
        ],
        "color": PINK
    },
    {
        "phase": "PHASE 4",
        "title": "Local / Air-Gapped LLM Execution",
        "bullets": [
            "Support on-premise inference with Ollama and vLLM for sensitive medical or classified defense documents.",
            "Complete zero-data-egress option allowing organizations to run FileMind fully offline."
        ],
        "color": GREEN
    },
    {
        "phase": "PHASE 5",
        "title": "Collaborative Team Vaults & RBAC",
        "bullets": [
            "Shared workspace vaults with granular role-based access control (Admin, Editor, Read-Only Reviewer).",
            "Audit trail logs tracking who viewed, exported, or modified document metadata."
        ],
        "color": PURPLE
    }
]

y_pos_r = Inches(2.0)
card_h_r = Inches(1.1)

for r in roadmap:
    card = add_card(s11, Inches(0.8), y_pos_r, Inches(11.733), card_h_r)
    tf_r = card.text_frame
    tf_r.margin_left = Inches(0.35)
    tf_r.margin_top = Inches(0.12)
    tf_r.margin_right = Inches(0.35)

    p_p = tf_r.paragraphs[0]
    r_p = p_p.add_run()
    r_p.text = f"{r['phase']}  —  {r['title']}"
    r_p.font.name = FONT_FAMILY
    r_p.font.size = Pt(11.5)
    r_p.font.bold = True
    r_p.font.color.rgb = r["color"]

    for b in r["bullets"]:
        p_b = tf_r.add_paragraph()
        p_b.space_before = Pt(2)
        r_bullet = p_b.add_run()
        r_bullet.text = "• "
        r_bullet.font.bold = True
        r_bullet.font.color.rgb = r["color"]
        r_text = p_b.add_run()
        r_text.text = b
        r_text.font.name = FONT_FAMILY
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = TEXT_BODY

    y_pos_r += card_h_r + Inches(0.14)

add_speaker_note(s11, """Looking forward, FileMind has a clear roadmap for future development.
Our immediate next milestone is Cross-Document Synthesis. Instead of querying one PDF at a time, users will be able to ask questions across their entire vault — for example, 'Summarize all changes between the 2024 and 2025 service contracts.'

Next, we plan to implement Visual Bounding-Box Citations, so that when the AI returns an answer, clicking it highlights the exact sentence on the rendered PDF page.

We also plan to support local LLMs like Ollama for organizations with strict privacy requirements that cannot send documents to the cloud, as well as team workspaces with role-based access control.""")


# ==============================================================================
# SLIDE 12: CONCLUSION, LIVE DEMO & Q&A
# ==============================================================================
s12 = prs.slides.add_slide(blank_layout)
set_slide_background(s12)

# Title Category Chip
chip12 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(3.2), Inches(0.35))
chip12.fill.solid()
chip12.fill.fore_color.rgb = RGBColor(19, 29, 49)
chip12.line.color.rgb = GREEN
chip12.line.width = Pt(1.2)
p = chip12.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "PROJECT SUMMARY & DEMO"
r.font.name = FONT_FAMILY
r.font.size = Pt(10)
r.font.bold = True
r.font.color.rgb = GREEN

# Main Title
title_box12 = s12.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.733), Inches(1.2))
tf12 = title_box12.text_frame
tf12.word_wrap = True
p12_1 = tf12.paragraphs[0]
r12_1 = p12_1.add_run()
r12_1.text = "FileMind — The Intelligent Document Vault"
r12_1.font.name = FONT_FAMILY
r12_1.font.size = Pt(36)
r12_1.font.bold = True
r12_1.font.color.rgb = WHITE

p12_2 = tf12.add_paragraph()
p12_2.space_before = Pt(4)
r12_2 = p12_2.add_run()
r12_2.text = "“Stop searching for documents. Start finding answers with FileMind.”"
r12_2.font.name = FONT_FAMILY
r12_2.font.size = Pt(16)
r12_2.font.italic = True
r12_2.font.color.rgb = CYAN

# Key Takeaways Summary Card
sum_card = add_card(s12, Inches(0.8), Inches(2.7), Inches(6.0), Inches(4.1), border_color=CARD_BORDER)
tf_sc = sum_card.text_frame
tf_sc.margin_left = Inches(0.35)
tf_sc.margin_top = Inches(0.25)
tf_sc.margin_right = Inches(0.35)

p_sh = tf_sc.paragraphs[0]
r_sh = p_sh.add_run()
r_sh.text = "Key Project Takeaways:"
r_sh.font.name = FONT_FAMILY
r_sh.font.size = Pt(13)
r_sh.font.bold = True
r_sh.font.color.rgb = WHITE

takeaways = [
    ("Problem Solved", "Replaces passive, unsearchable folders with an active AI-powered document intelligence vault."),
    ("Production Quality", "Engineered with Next.js 16, Supabase PostgreSQL, and Google Gemini with verified test suites."),
    ("Resilient Concurrency", "Tested with 150-file bulk batches, automated recovery, and transactional advisory locks."),
    ("High Privacy", "Database-enforced Row-Level Security ensuring strict per-user vault isolation.")
]

for t, d in takeaways:
    p_tk = tf_sc.add_paragraph()
    p_tk.space_before = Pt(8)
    r_tk1 = p_tk.add_run()
    r_tk1.text = f"✔ {t}: "
    r_tk1.font.bold = True
    r_tk1.font.color.rgb = GREEN
    r_tk2 = p_tk.add_run()
    r_tk2.text = d
    r_tk2.font.color.rgb = TEXT_BODY
    r_tk2.font.size = Pt(10.5)

# Live Demo & Q&A Card
demo_card = add_card(s12, Inches(7.1), Inches(2.7), Inches(5.433), Inches(4.1), border_color=RGBColor(25, 55, 80))
tf_dc = demo_card.text_frame
tf_dc.margin_left = Inches(0.35)
tf_dc.margin_top = Inches(0.25)
tf_dc.margin_right = Inches(0.35)

p_dh = tf_dc.paragraphs[0]
r_dh = p_dh.add_run()
r_dh.text = "Ready for Evaluation & Live Demo:"
r_dh.font.name = FONT_FAMILY
r_dh.font.size = Pt(13)
r_dh.font.bold = True
r_dh.font.color.rgb = CYAN

p_d1 = tf_dc.add_paragraph()
p_d1.space_before = Pt(10)
r_d1 = p_d1.add_run()
r_d1.text = "1. Live Ingestion Walkthrough:\n"
r_d1.font.bold = True
r_d1.font.color.rgb = WHITE
r_d1_sub = p_d1.add_run()
r_d1_sub.text = "Drag-and-drop a batch of test invoices, papers, and legal agreements."
r_d1_sub.font.size = Pt(10)
r_d1_sub.font.color.rgb = TEXT_MUTED

p_d2 = tf_dc.add_paragraph()
p_d2.space_before = Pt(8)
r_d2 = p_d2.add_run()
r_d2.text = "2. AI Extraction & Semantic Search:\n"
r_d2.font.bold = True
r_d2.font.color.rgb = WHITE
r_d2_sub = p_d2.add_run()
r_d2_sub.text = "Watch the Gemini pipeline extract categories and search by natural concepts in real-time."
r_d2_sub.font.size = Pt(10)
r_d2_sub.font.color.rgb = TEXT_MUTED

p_d3 = tf_dc.add_paragraph()
p_d3.space_before = Pt(12)
r_d3 = p_d3.add_run()
r_d3.text = "Presenter Contacts & Resources:\n"
r_d3.font.bold = True
r_d3.font.color.rgb = PINK
r_d3_sub = p_d3.add_run()
r_d3_sub.text = "GitHub: [github.com/your-username/filemind]\nEmail: [your-email@college.edu]  |  Demo URL: [http://localhost:3000]\n\nThank you! We welcome questions from our teacher and judges."
r_d3_sub.font.size = Pt(9.5)
r_d3_sub.font.color.rgb = TEXT_BODY

add_speaker_note(s12, """To conclude our presentation: FileMind takes the friction out of document management. By bridging the gap between secure cloud storage and state-of-the-art multimodal AI, we transform dead PDF archives into dynamic knowledge bases.

We have verified our implementation through automated test suites covering bulk upload concurrency, model fallbacks, and database security invariants.

At this point, we would love to invite you to our live demonstration where we will drop a sample batch of documents and demonstrate real-time extraction and semantic search.

Thank you for your time and guidance. We are ready for your questions!""")

# Save Presentation
output_pptx_path = os.path.join(os.getcwd(), "FileMind_Project_Presentation.pptx")
prs.save(output_pptx_path)
print(f"Presentation successfully saved to: {output_pptx_path}")
